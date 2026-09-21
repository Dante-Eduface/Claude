"""Controleert de gebouwde pptx op het bestand zelf, niet op de preview.

Waarom dit bestaat: de HTML-preview liet afbeeldingen netjes in hun vak passen
(object-fit: contain) terwijl pptx ze oprekt. Daardoor bleven een 14% uitgerekt
Eduface-logo en voor 30% afgesneden iconen vijftien slides lang onzichtbaar voor
elke controle die ik deed.

Dat is nu op twee manieren dichtgezet:
  1. render_html gebruikt 'fill', net als pptx, dus de preview kan niet meer
     mooier zijn dan het echte bestand;
  2. dit script meet de geleverde pptx na en weigert groen licht te geven als
     een afbeelding vervormd staat of buiten de veilige marge valt.

LibreOffice zou de pptx echt kunnen renderen, maar is in deze omgeving stuk
(`soffice --convert-to pdf` faalt zelfs op een tekstbestand), dus visuele
controle loopt via de previews, die nu wel eerlijk zijn.

  python3 check_pptx.py
"""
import io
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
os.chdir(HERE)

PPTX = 'UTI-Eduface-deck.pptx'
PX = 6350                      # EMU per ontwerp-pixel
W, H, MARGE = 1920, 1080, 120
REFERENTIE_PDF = ('/root/.claude/uploads/0449f0ab-9edd-598c-b3bd-ed4372c7b86b/'
                  'c8727ba2-UTI_-_Eduface_presentation.pdf')

AANTAL_SLIDES = 16

MOET_ERIN = [
    (2, 'agenda'),
    (2, 'demo'),
    (4, '2 - 4 hrs'),
    (4, 'tutoring and remedial support'),
    (5, 'new campuses per year'),
    (7, 'largest variable cost items'),
    (7, 'higher course completion rates'),
    (12, 'blackboard and canvas integration is a precondition'),
    (13, 'demo'),
    (16, 'programmatic expansion'),
    (16, 'target launch december 2026'),
    # 'legal and it validation' stond bij Validate in the U.S.; die stap is op
    # verzoek van Dante van de slide gehaald, dus die tekst hoort hier niet meer.
]


def norm(s):
    s = s.replace('’', "'").replace('–', '-').replace('—', '-')
    return re.sub(r'\s+', ' ', s).strip().lower()


def controleer():
    from pptx import Presentation
    from PIL import Image

    p = Presentation(PPTX)
    fouten = []

    n_slides = len(p.slides._sldIdLst)
    if n_slides != AANTAL_SLIDES:
        fouten.append(f'{n_slides} slides in plaats van {AANTAL_SLIDES}')

    breedte_inch = round(p.slide_width / 914400, 2)
    if breedte_inch != 13.33:
        fouten.append(f'slidebreedte {breedte_inch} inch in plaats van 13.33')

    for n, s in enumerate(p.slides, 1):
        tekstvakken = 0
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                tekstvakken += 1

            if sh.shape_type == 13:                      # afbeelding
                im = Image.open(io.BytesIO(sh.image.blob))
                na = im.width / im.height
                pa = sh.width / sh.height
                if abs(pa - na) / na > 0.02:
                    fouten.append(f'slide {n}: afbeelding {im.size} staat als '
                                  f'{pa:.2f} maar hoort {na:.2f} te zijn')

            # niets mag buiten de veilige marge vallen, behalve het volvlaks
            # achtergrondvlak dat de assembler onderop legt
            x, y = sh.left / PX, sh.top / PX
            br, ho = sh.width / PX, sh.height / PX
            volvlaks = br >= W - 1 and ho >= H - 1
            if not volvlaks and (x < MARGE - 1 or y < 0 or
                                 x + br > W - MARGE + 1 or y + ho > H):
                fouten.append(f'slide {n}: vorm buiten de marge op '
                              f'{x:.0f},{y:.0f} ({br:.0f}x{ho:.0f})')

        if tekstvakken == 0:
            fouten.append(f'slide {n}: geen bewerkbare tekst')

    # verplichte tekstfragmenten: Dantes handmatige wijzigingen plus de inhoud
    # van de vorige ronde, zodat er niets stilletjes uit kan vallen
    for n, frag in MOET_ERIN:
        blad = norm(' '.join(sh.text_frame.text for sh in p.slides[n - 1].shapes
                             if sh.has_text_frame))
        if frag not in blad:
            fouten.append(f'slide {n}: tekst ontbreekt: "{frag}"')

    return fouten


if __name__ == '__main__':
    fouten = controleer()
    if fouten:
        print(f'{len(fouten)} probleem(en):')
        for f in fouten:
            print('  -', f)
        sys.exit(1)
    print(f'{PPTX} in orde: {AANTAL_SLIDES} slides, geen vervormde afbeeldingen, '
          'niets buiten de marge, alle verplichte teksten aanwezig')
