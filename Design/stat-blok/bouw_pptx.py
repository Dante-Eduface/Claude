"""Variant C als bewerkbare pptx, op de maat van het blok zelf (voor Canva)."""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

PX = 6350                                   # 1 ontwerp-pixel = een halve punt
W, H = 1624, 354
NAVY, INK50, INK500, INK700, WHITE = '002333', 'F3F7F8', '5B7480', '2C4A57', 'FFFFFF'
HEAD, BODY = 'League Spartan', 'Inter'

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(W * PX), Emu(H * PX)
s = prs.slides.add_slide(prs.slide_layouts[6])

def kaart(x, fill):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x * PX), 0, Emu(764 * PX), Emu(H * PX))
    sh.adjustments[0] = 20 / H
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    sh.line.fill.background(); sh.shadow.inherit = False

def tekst(x, y, w, h, t, font, size, color, bold=False, spc=None, ls=None):
    tb = s.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    if ls: p.line_spacing = ls
    r = p.add_run(); r.text = t
    r.font.name = font; r.font.size = Pt(size / 2); r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    if spc: r._r.get_or_add_rPr().set('spc', str(spc))

def inhoud(x, label, getal, sub, donker):
    tekst(x + 64, 56, 636, 30, label, BODY, 22, '9EABB1' if donker else INK500, bold=True, spc=88)
    tekst(x + 64, 96, 636, 140, getal, HEAD, 140, WHITE if donker else NAVY, bold=True, ls=0.8)
    tekst(x + 64, 263, 636, 40, sub, BODY, 26, WHITE if donker else INK700)

kaart(0, INK50)
inhoud(0, 'NU', '€ 37.000', 'Wat nakijken de opleiding nu kost', False)
kaart(860, NAVY)
inhoud(860, 'MET EDUFACE', '106%', 'Rendement op de investering', True)

pijl = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(784 * PX), Emu(177 * PX), Emu(840 * PX), Emu(177 * PX))
pijl.line.color.rgb = RGBColor.from_string(NAVY); pijl.line.width = Emu(3 * PX)
ln = pijl.line._get_or_add_ln()
tail = ln.makeelement(qn('a:tailEnd'), {'type': 'arrow', 'w': 'med', 'len': 'med'}); ln.append(tail)

prs.save('stat-blok-C-bewerkbaar.pptx')
