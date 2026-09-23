"""Maakt uit elk criteria-bestand een leeg beoordelingsformulier (pdf plus docx).

Draaien: python3 build.py
De pdf komt uit headless Chromium en de docx uit python-docx, beide uit dezelfde
criteria, dus ze kunnen niet uiteen lopen. Een formulier erbij? Zet een regel in
FORMULIEREN en schrijf het bijbehorende criteria-bestand.
"""
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HIER = Path(__file__).parent
CHROMIUM = "/opt/pw-browsers/chromium"
TEKSTBREEDTE_CM = 17.0

FORMULIEREN = [
    {
        "bron": "criteria-cat-diagnostiek.md",
        "basis": "beoordelingsformulier-cat-diagnostiek",
        "titel": "Beoordelingsformulier CAT diagnostiek",
        "ondertitel": None,
        "kopvelden": ["Datum", "Naam student", "Beoordelaar"],
        "oordeelkop": "Voldaan / Niet voldaan",
        "oordeelbreedte": 0.22,
        "rijhoogte_cm": None,
        "slotvelden": ["Beoordeling", "Feedback"],
    },
    {
        "bron": "criteria-cat-mkf.md",
        "basis": "beoordelingsformulier-cat-mkf",
        "titel": "Beoordelingsformulier CAT Master Kinderfysiotherapie",
        "ondertitel": "MKF26-WO1-CAT-1-5 · Breederode Hogeschool",
        "kopvelden": ["Groep: MKF26", "Namen", "Titel", "Datum"],
        "oordeelkop": "Feedback",
        "oordeelbreedte": 0.42,
        "rijhoogte_cm": 1.1,
        "slotvelden": ["Eindoordeel (voldaan / niet voldaan)"],
    },
]


def lees_criteria(pad):
    """Geeft [(sectietitel, [criterium, ...]), ...] terug."""
    secties = []
    for regel in pad.read_text(encoding="utf-8").splitlines():
        regel = regel.rstrip()
        if regel.startswith("## "):
            secties.append((regel[3:].strip(), []))
        elif regel.startswith("- ") and secties:
            secties[-1][1].append(regel[2:].strip())
    return secties


def splits_veld(veld):
    """'Groep: MKF26' wordt ('Groep', 'MKF26'), 'Datum' wordt ('Datum', '')."""
    label, _, waarde = veld.partition(":")
    return label.strip(), waarde.strip()


# ---------------------------------------------------------------- docx


def arceer(cel, hex_kleur):
    vulling = OxmlElement("w:shd")
    vulling.set(qn("w:val"), "clear")
    vulling.set(qn("w:fill"), hex_kleur)
    cel._tc.get_or_add_tcPr().append(vulling)


def zet_hoogte(rij, cm):
    hoogte = OxmlElement("w:trHeight")
    hoogte.set(qn("w:val"), str(int(cm * 567)))
    hoogte.set(qn("w:hRule"), "atLeast")
    rij._tr.get_or_add_trPr().append(hoogte)


def zet_tekst(cel, tekst, vet=False, grootte=10, midden=False):
    alinea = cel.paragraphs[0]
    alinea.paragraph_format.space_before = Pt(2)
    alinea.paragraph_format.space_after = Pt(2)
    if midden:
        alinea.alignment = WD_ALIGN_PARAGRAPH.CENTER
    stuk = alinea.add_run(tekst)
    stuk.bold = vet
    stuk.font.size = Pt(grootte)
    stuk.font.name = "Calibri"
    cel.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def bouw_docx(vorm, secties, pad):
    doc = Document()
    for sectie in doc.sections:
        sectie.top_margin = Cm(2)
        sectie.bottom_margin = Cm(2)
        sectie.left_margin = Cm(2)
        sectie.right_margin = Cm(2)

    standaard = doc.styles["Normal"]
    standaard.font.name = "Calibri"
    standaard.font.size = Pt(11)

    kop = doc.add_paragraph()
    kop.paragraph_format.space_after = Pt(2 if vorm["ondertitel"] else 10)
    stuk = kop.add_run(vorm["titel"])
    stuk.bold = True
    stuk.font.size = Pt(16)

    if vorm["ondertitel"]:
        onder = doc.add_paragraph()
        onder.paragraph_format.space_after = Pt(10)
        stuk = onder.add_run(vorm["ondertitel"])
        stuk.font.size = Pt(10)

    for veld in vorm["kopvelden"]:
        label, waarde = splits_veld(veld)
        alinea = doc.add_paragraph()
        alinea.paragraph_format.space_after = Pt(2)
        alinea.add_run(f"{label}: ").bold = True
        if waarde:
            alinea.add_run(waarde)

    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    tabel = doc.add_table(rows=1, cols=2)
    tabel.style = "Table Grid"
    tabel.autofit = False
    oordeel_cm = TEKSTBREEDTE_CM * vorm["oordeelbreedte"]
    breedtes = (Cm(TEKSTBREEDTE_CM - oordeel_cm), Cm(oordeel_cm))

    koprij = tabel.rows[0]
    zet_tekst(koprij.cells[0], "Beoordelingscriteria", vet=True)
    zet_tekst(koprij.cells[1], vorm["oordeelkop"], vet=True, midden=True)
    for cel in koprij.cells:
        arceer(cel, "D9D9D9")

    for sectietitel, criteria in secties:
        rij = tabel.add_row()
        cel = rij.cells[0].merge(rij.cells[1])
        zet_tekst(cel, sectietitel, vet=True)
        arceer(cel, "F2F2F2")
        for criterium in criteria:
            rij = tabel.add_row()
            zet_tekst(rij.cells[0], criterium)
            zet_tekst(rij.cells[1], "")
            if vorm["rijhoogte_cm"]:
                zet_hoogte(rij, vorm["rijhoogte_cm"])

    for rij in tabel.rows:
        for cel, breedte in zip(rij.cells, breedtes):
            cel.width = breedte

    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    for veld in vorm["slotvelden"]:
        alinea = doc.add_paragraph()
        alinea.paragraph_format.space_after = Pt(18)
        alinea.add_run(f"{veld}: ").bold = True

    doc.save(pad)
    return pad


# ----------------------------------------------------------------- pdf

HTML_SJABLOON = """<!doctype html>
<html lang="nl"><head><meta charset="utf-8"><title>{titel}</title>
<style>
  @page {{ size: A4; margin: 20mm; }}
  body {{ font-family: Calibri, Carlito, Arial, sans-serif; font-size: 10pt; color: #000; margin: 0; }}
  h1 {{ font-size: 16pt; margin: 0 0 4px; }}
  .ondertitel {{ font-size: 10pt; margin: 0 0 12px; }}
  .veld {{ font-size: 11pt; margin-bottom: 3px; }}
  table {{ width: 100%; border-collapse: collapse; margin-top: 14px; table-layout: fixed; }}
  th, td {{ border: 1px solid #000; padding: 4px 6px; vertical-align: middle; text-align: left; }}
  th {{ background: #d9d9d9; font-weight: 700; }}
  td.oordeel, th.oordeel {{ width: {oordeelbreedte}%; text-align: center; }}
  td.oordeel {{ height: {rijhoogte}mm; }}
  tr.sectie td {{ background: #f2f2f2; font-weight: 700; }}
  tr.sectie {{ break-after: avoid; page-break-after: avoid; }}
  tr {{ page-break-inside: avoid; }}
  .slot {{ margin-top: 18px; font-size: 11pt; }}
  .slot p {{ margin: 0 0 24px; }}
</style></head><body>
<h1>{titel}</h1>
{ondertitel}
{velden}
<table>
<tr><th>Beoordelingscriteria</th><th class="oordeel">{oordeelkop}</th></tr>
{rijen}
</table>
<div class="slot">{slot}</div>
</body></html>
"""


def ontsnap(tekst):
    return tekst.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def maak_html(vorm, secties):
    velden = []
    for veld in vorm["kopvelden"]:
        label, waarde = splits_veld(veld)
        staart = f" {ontsnap(waarde)}" if waarde else ""
        velden.append(f'<div class="veld"><b>{ontsnap(label)}:</b>{staart}</div>')

    rijen = []
    for sectietitel, criteria in secties:
        rijen.append(f'<tr class="sectie"><td colspan="2">{ontsnap(sectietitel)}</td></tr>')
        for criterium in criteria:
            rijen.append(f'<tr><td>{ontsnap(criterium)}</td><td class="oordeel"></td></tr>')

    ondertitel = ""
    if vorm["ondertitel"]:
        ondertitel = f'<div class="ondertitel">{ontsnap(vorm["ondertitel"])}</div>'

    return HTML_SJABLOON.format(
        titel=ontsnap(vorm["titel"]),
        ondertitel=ondertitel,
        velden="\n".join(velden),
        oordeelkop=ontsnap(vorm["oordeelkop"]),
        oordeelbreedte=round(vorm["oordeelbreedte"] * 100),
        rijhoogte=round((vorm["rijhoogte_cm"] or 0.6) * 10),
        rijen="\n".join(rijen),
        slot="\n".join(f"<p><b>{ontsnap(v)}:</b></p>" for v in vorm["slotvelden"]),
    )


def bouw_pdf(vorm, secties, pad):
    pad_html = pad.with_suffix(".html")
    pad_html.write_text(maak_html(vorm, secties), encoding="utf-8")
    try:
        subprocess.run(
            [CHROMIUM, "--headless", "--disable-gpu", "--no-sandbox",
             "--no-pdf-header-footer", f"--print-to-pdf={pad}", pad_html.as_uri()],
            check=True, capture_output=True, timeout=180,
        )
    except Exception as fout:
        print(f"geen pdf gemaakt: {fout}")
        return None
    finally:
        pad_html.unlink(missing_ok=True)
    return pad


if __name__ == "__main__":
    for vorm in FORMULIEREN:
        secties = lees_criteria(HIER / vorm["bron"])
        aantal = sum(len(c) for _, c in secties)
        pdf = bouw_pdf(vorm, secties, HIER / f"{vorm['basis']}.pdf")
        docx = bouw_docx(vorm, secties, HIER / f"{vorm['basis']}.docx")
        print(f"{vorm['basis']}: {len(secties)} secties, {aantal} criteria"
              f" -> {pdf.name if pdf else 'geen pdf'}, {docx.name}")
