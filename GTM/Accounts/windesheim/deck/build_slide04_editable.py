"""Slide 4 als bewerkbare pptx: echte tekstvakken en vormen, geen platte afbeelding.
Ontwerpraster is 1920x1080 px; 1 px = 6350 EMU op een 16:9 canvas van 13,333 inch."""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PX = 6350
def E(px): return Emu(int(round(px * PX)))
def P(px): return Pt(px / 2)          # 1920 px over 13,333 inch -> px/2 = pt

NAVY   = RGBColor(0x00, 0x23, 0x33)
AMBER  = RGBColor(0xE0, 0x7B, 0x00)
TINT   = RGBColor(0xFD, 0xF0, 0xDF)
INK500 = RGBColor(0x5B, 0x74, 0x80)
INK200 = RGBColor(0xCD, 0xD9, 0xDE)
INK100 = RGBColor(0xE7, 0xEE, 0xF0)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

HEAD, BODY = 'League Spartan', 'Inter'

prs = Presentation()
prs.slide_width, prs.slide_height = E(1920), E(1080)
s = prs.slides.add_slide(prs.slide_layouts[6])

def text(x, y, w, h, txt, *, font=BODY, size=32, bold=False, color=NAVY,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    r = p.add_run(); r.text = txt
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = font, P(size), bold, color
    return tb

def rect(x, y, w, h, fill):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh

def rrect(x, y, w, h, fill, radius):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background(); sh.shadow.inherit = False
    sh.adjustments[0] = radius / min(w, h)
    return sh

def oval(x, y, d, fill, line, lw):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, E(x), E(y), E(d), E(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line; sh.line.width = P(lw); sh.shadow.inherit = False
    return sh

# ---- titel ---------------------------------------------------------------
text(120, 88, 1680, 120, 'Visitatie', font=HEAD, size=96, bold=True, spacing=1.0)

# ---- tijdlijn ------------------------------------------------------------
rect(120, 428, 1680, 4, INK100)

COLS = [
    ('2021', 'Verschillen',        'assets/icon-verschillen.png'),
    ('2022', 'Variatie',           'assets/icon-variatie.png'),
    ('2024', 'Wisselend ingevuld', 'assets/icon-wisselend.png'),
    ('2025', 'Cijfer ≠ feedback', 'assets/icon-ongelijk.png'),
]
COLW, GAP = 366, 72
for i, (yr, lbl, icon) in enumerate(COLS):
    cx = 120 + i * (COLW + GAP) + COLW / 2
    text(cx - 100, 336, 200, 48, yr, font=HEAD, size=34, bold=True, color=INK500, spacing=1.0)
    oval(cx - 12, 418, 24, WHITE, NAVY, 6)
    rect(cx - 2, 442, 4, 56, INK200)
    rrect(cx - 74, 506, 148, 148, TINT, 36)
    s.shapes.add_picture(icon, E(cx - 37), E(543), E(74), E(74))
    text(cx - COLW / 2, 680, COLW, 130, lbl, font=HEAD, size=46, bold=True, spacing=1.1)

# ---- voet ----------------------------------------------------------------
rect(120, 868, 1680, 1, INK100)
text(120, 892, 1320, 56,
     '“De narratieve feedback is niet altijd helemaal passend bij de rubric.”',
     size=38, align=PP_ALIGN.LEFT, spacing=1.3)
text(120, 950, 1320, 64,
     'NQA, Ad E-commerce Windesheim, visitatie november 2024, p. 24. '
     'Zelfde patroon in vier andere visitatierapporten, 2021-2025.',
     size=23, color=INK500, align=PP_ALIGN.LEFT, spacing=1.3)

# ---- logo's --------------------------------------------------------------
s.shapes.add_picture('logo_navy.png', E(1419), E(907), height=E(46))
rect(1629, 904, 1, 52, INK200)
s.shapes.add_picture('assets/windesheim.png', E(1670), E(900), height=E(60))

text(1700, 1004, 100, 36, '4', size=23, color=INK200, align=PP_ALIGN.RIGHT)

prs.save('slide-04-visitatie-bewerkbaar.pptx')
print('geschreven: slide-04-visitatie-bewerkbaar.pptx')
