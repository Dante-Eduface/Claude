"""Bouwt het wekelijkse Sales update-deck voor de maandagsessie met Tjarko.

Gebruik (vanuit deze map):
    python3 bouw.py 2026-09-28               heel deck: pptx plus een PNG per slide
    python3 bouw.py 2026-09-28 --slides 1,2  alleen deze slides als PNG, geen pptx

Leest <datum>/invoer.json. De mapnaam is de datum van de meeting (een maandag).
Datum, weeknummer en de dagen tot einde kwartaal rekent dit script zelf uit, de rest
levert Dante aan: het aantal gestarte salesprocessen, de uitdaging en de losse punten.

Bouwt op de pitch-deck-pipeline (.claude/skills/pitch-deck/deckbuild.py) zonder die aan
te passen. Maten en kleuren komen uit Design/System/slides/tokens.css en core/tokens.css.
"""
import calendar, datetime as dt, glob, json, os, sys, tempfile

HIER = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HIER, '..', '..'))
sys.path.insert(0, os.path.join(ROOT, '.claude', 'skills', 'pitch-deck'))
sys.dont_write_bytecode = True     # geen __pycache__ in de skillmap

import deckbuild as D
from deckbuild import rect, pic, text, render_html, render_pptx
from deckbuild import NAVY, WHITE, INK100, INK500, HEAD, BODY
from PIL import Image, ImageFont

# deckbuild wijst naar Chrome op de Mac. Staat die er niet, dan de Playwright-Chromium.
# De headless shell houdt het venster op precies 1920x1080; gewone headless Chrome
# snoept er ruimte af voor de browserbalk en laat onderaan een witte strook.
if not os.path.exists(D.CHROME):
    kandidaten = (sorted(glob.glob('/opt/pw-browsers/chromium_headless_shell-*/chrome-linux/headless_shell'))
                  or sorted(glob.glob('/opt/pw-browsers/chromium-*/chrome-linux/chrome')))
    if kandidaten:
        D.CHROME = kandidaten[-1]



def shoot(html_path, png_path):
    """Zoals deckbuild.shoot, maar met --no-sandbox als root (cloudcontainer), en met foutmelding."""
    import subprocess
    args = [D.CHROME, '--headless', '--disable-gpu', f'--screenshot={png_path}',
            '--window-size=1920,1080', '--hide-scrollbars', '--virtual-time-budget=4000']
    if hasattr(os, 'geteuid') and os.geteuid() == 0:
        args.append('--no-sandbox')
    r = subprocess.run(args + [os.path.abspath(html_path)], capture_output=True, text=True)
    if not os.path.exists(png_path):
        sys.exit(f'Screenshot mislukt voor {html_path}:\n{r.stderr[-800:]}')
    return png_path


# ---------------------------------------------------------------- tokens
M, W, H = 120, 1920, 1080
BREED = W - 2 * M
HERO, TITEL, SUB, TEKST, BIJSCHRIFT, STAT = 132, 76, 44, 32, 24, 180
ZACHT_INV = '9EABB1'      # wit op 62% over navy (foreground-inverse-soft), pptx kent geen alfa
DOEL_MAAND = 20           # Context/current-priorities.md

MERK = os.path.join(ROOT, 'Design', 'Merk')
LOGO_NAVY = os.path.join(MERK, 'logos', 'logo_navy.png')
KOP_FONT = os.path.join(MERK, 'fonts', 'LeagueSpartan-Bold.ttf')
TEKST_FONT = os.path.join(MERK, 'fonts', 'Inter-Regular.ttf')

MAANDEN = ['januari', 'februari', 'maart', 'april', 'mei', 'juni', 'juli',
           'augustus', 'september', 'oktober', 'november', 'december']
DAGEN = ['maandag', 'dinsdag', 'woensdag', 'donderdag', 'vrijdag', 'zaterdag', 'zondag']


def logo_wit():
    """Het witte logo staat met veel lege rand in een vierkant. Knip die weg, buiten de repo."""
    uit = os.path.join(tempfile.gettempdir(), 'eduface-logo-wit-bijgesneden.png')
    if not os.path.exists(uit):
        im = Image.open(os.path.join(MERK, 'logos', 'logo_white.png')).convert('RGBA')
        masker = im.split()[3].point(lambda v: 255 if v > 20 else 0)
        im.crop(masker.getbbox()).save(uit)
    return uit


# ---------------------------------------------------------------- meten
_fonts = {}


def breedte(t, size, kop=False):
    # Beide fonts zijn variabel. League Spartan staat standaard op Thin, dus zet hem op Bold.
    if (kop, size) not in _fonts:
        f = ImageFont.truetype(KOP_FONT if kop else TEKST_FONT, size)
        if kop:
            f.set_variation_by_name('Bold')
        _fonts[kop, size] = f
    return _fonts[kop, size].getlength(t)


def regels(t, size, max_w, kop=False):
    """Hoeveel regels wordt deze tekst in een kader van max_w breed."""
    n, regel = 1, ''
    for woord in t.split():
        proef = (regel + ' ' + woord).strip()
        if breedte(proef, size, kop) > max_w and regel:
            n, regel = n + 1, woord
        else:
            regel = proef
    return n


# ---------------------------------------------------------------- data
def kalender(maandag):
    """Alles wat uit de datum van de meeting volgt."""
    zondag = maandag - dt.timedelta(days=1)            # laatste dag van de telling
    week_start = maandag - dt.timedelta(days=7)
    # Aftellen naar het einde van het kwartaal waarin de zondag van déze week valt.
    # Op maandag 28-09 is dat Q4 en niet "nog 2 dagen Q3" (besluit 27-09-2026).
    eind_week = maandag + dt.timedelta(days=6)
    q = (eind_week.month - 1) // 3 + 1
    kw_maand = 3 * q
    kw_eind = dt.date(eind_week.year, kw_maand, calendar.monthrange(eind_week.year, kw_maand)[1])
    dagen_in_maand = calendar.monthrange(zondag.year, zondag.month)[1]
    return dict(
        maandag=maandag, zondag=zondag, week_start=week_start,
        week=maandag.isocalendar()[1], kwartaal=q,
        dagen_kwartaal=(kw_eind - maandag).days,
        maand=MAANDEN[zondag.month - 1],
        tempo=round(DOEL_MAAND * zondag.day / dagen_in_maand),
    )


def datum_lang(d):
    return f'{DAGEN[d.weekday()].capitalize()} {d.day} {MAANDEN[d.month - 1]} {d.year}'


def periode(a, b):
    if a.month == b.month:
        return f'{a.day} t/m {b.day} {MAANDEN[b.month - 1]}'
    return f'{a.day} {MAANDEN[a.month - 1]} t/m {b.day} {MAANDEN[b.month - 1]}'


# ---------------------------------------------------------------- bouwstenen
def navy_vlak():
    # render_pptx schildert geen achtergrond, dus een navy slide heeft een vlak nodig.
    return [rect(0, 0, W, H, NAVY)]


def kop(t, dark=False):
    return [text(M, 88, BREED, 90, t, font=HEAD, size=TITEL, bold=True,
                 color=WHITE if dark else NAVY, align='left', ls=1.04)]


def voet(bron=None, dark=False):
    out = []
    if bron:
        out.append(text(M, 922, 1300, 36, bron, size=BIJSCHRIFT,
                        color=ZACHT_INV if dark else INK500, align='left', ls=1.35))
    out.append(pic(logo_wit() if dark else LOGO_NAVY, W - M - 170, 914, 170, 46))
    return out


# ---------------------------------------------------------------- slides
def slide_titel(inv, k):
    els = navy_vlak()
    els += [text(M, 330, BREED, 140, 'Sales update', font=HEAD, size=HERO, bold=True,
                 color=WHITE, align='left', ls=0.98),
            text(M, 492, BREED, 56, f"{datum_lang(k['maandag'])}, week {k['week']}",
                 size=SUB, color=ZACHT_INV, align='left', ls=1.15)]
    # Aftellen rechtsonder, logo linksonder (layouts.md: opening).
    els += [text(W - M - 600, 812, 600, 84, str(k['dagen_kwartaal']), font=HEAD, size=TITEL,
                 bold=True, color=WHITE, align='right', ls=1.04),
            text(W - M - 600, 904, 600, 44, f"dagen tot einde Q{k['kwartaal']}",
                 size=TEKST, color=ZACHT_INV, align='right', ls=1.35),
            pic(logo_wit(), M, 910, 170, 46)]
    return els, True


def slide_gestart(inv, k):
    els = kop('Salesprocessen gestart')
    maand_n, week_n = inv['gestart_maand'], inv['gestart_week']

    # Links de maand tegen het doel, rechts de week. Ongelijk gewicht: de maand is de maat.
    lx, lw = M, 1040
    rx = lx + lw + 128
    rw = W - M - rx
    y_label, y_getal, y_balk = 348, 402, 642

    els += [text(lx, y_label, lw, 44, inv.get('maand_label', k['maand']).capitalize(),
                 size=TEKST, color=INK500, align='left'),
            text(lx, y_getal, 600, 170, str(maand_n), font=HEAD, size=STAT, bold=True,
                 align='left', ls=0.9)]
    van_x = lx + breedte(str(maand_n), STAT, kop=True) + 28
    els += [text(van_x, y_getal + 88, 500, 56, f'van {DOEL_MAAND}', font=HEAD, size=SUB,
                 bold=True, color=INK500, align='left', ls=1.15)]

    # Balk naar 20, met een streep waar je vandaag naar verhouding zou staan.
    # Navy en grijs, nooit groen: kleur codeert geen oordeel (core/regels.md).
    els += [rect(lx, y_balk, lw, 24, INK100, r=12)]
    if maand_n > 0:
        els += [rect(lx, y_balk, max(24, lw * min(maand_n, DOEL_MAAND) / DOEL_MAAND), 24, NAVY, r=12)]
    sx = lx + lw * k['tempo'] / DOEL_MAAND
    els += [rect(sx - 2, y_balk - 20, 4, 64, INK500)]
    lab = f"Tempo voor {DOEL_MAAND}: {k['tempo']} op {k['zondag'].day} {k['maand'][:3]}"
    lab_w = 440
    lab_x = min(max(sx - lab_w / 2, lx), lx + lw - lab_w)
    els += [text(lab_x, y_balk + 58, lab_w, 36, lab, size=BIJSCHRIFT, color=INK500, ls=1.35)]

    els += [text(rx, y_label, rw, 44, 'Vorige week', size=TEKST, color=INK500, align='left'),
            text(rx, y_getal, rw, 170, str(week_n), font=HEAD, size=STAT, bold=True,
                 align='left', ls=0.9),
            text(rx, y_balk - 10, rw, 44, periode(k['week_start'], k['zondag']),
                 size=TEKST, color=INK500, align='left')]

    bron = f"Telling: Dante, t/m {DAGEN[6]} {k['zondag'].day} {k['maand']}."
    if inv.get('voorbeeld'):
        bron = 'Voorbeeldcijfers, nog niet de echte telling. ' + bron
    els += voet(bron)
    return els, False


def slide_funnel(inv, k):
    return kop('Funnel overview') + voet(), False


def slide_uitdaging(inv, k):
    els = navy_vlak()
    zin, toel = inv['uitdaging'], inv.get('uitdaging_toelichting', '')
    size = HERO
    if regels(zin, HERO, BREED, kop=True) > 3:
        size = TITEL          # past het niet op hero-maat, dan de kopmaat
    n = regels(zin, size, BREED, kop=True)
    ls = 0.98 if size == HERO else 1.04
    zin_h = n * size * ls
    toel_n = regels(toel, TEKST, 1300) if toel else 0
    blok_h = 44 + 32 + zin_h + (48 + toel_n * TEKST * 1.45 if toel else 0)
    y = max(150, (868 - blok_h) / 2 + 40)
    els += [text(M, y, BREED, 44, 'Grootste uitdaging', size=TEKST, color=ZACHT_INV, align='left')]
    y += 44 + 32
    els += [text(M, y, BREED, zin_h + 10, zin, font=HEAD, size=size, bold=True,
                 color=WHITE, align='left', ls=ls)]
    if toel:
        y += zin_h + 48
        els += [text(M, y, 1300, toel_n * TEKST * 1.45 + 10, toel, size=TEKST,
                     color=ZACHT_INV, align='left', ls=1.45)]
    els += voet(dark=True)
    return els, True


def slides_punten(inv, k):
    """Maximaal zes punten per slide. Meer, dan een tweede slide."""
    punten = inv.get('losse_punten', [])
    groepen = [punten[i:i + 6] for i in range(0, len(punten), 6)] or [[]]
    uit = []
    for g, groep in enumerate(groepen):
        els = kop('Losse punten' + (f' ({g + 1})' if len(groepen) > 1 else ''))
        size, ls, tx = SUB, 1.15, M + 56
        y = 250
        for p in groep:
            n = regels(p, size, W - M - tx)
            els += [rect(M, y + 16, 16, 16, NAVY, r=8),
                    text(tx, y, W - M - tx, n * size * ls + 6, p, size=size, align='left', ls=ls)]
            y += n * size * ls + 40
        els += voet()
        uit.append((els, False))
    return uit


# ---------------------------------------------------------------- bouwen
def deck(inv, k):
    s = [slide_titel(inv, k), slide_gestart(inv, k), slide_funnel(inv, k), slide_uitdaging(inv, k)]
    return s + slides_punten(inv, k)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    datum = sys.argv[1]
    keuze = None
    if '--slides' in sys.argv:
        keuze = {int(x) for x in sys.argv[sys.argv.index('--slides') + 1].split(',')}

    map_ = os.path.join(HIER, datum)
    inv = json.load(open(os.path.join(map_, 'invoer.json'), encoding='utf-8'))
    maandag = dt.date.fromisoformat(datum)
    if maandag.weekday() != 0:
        print(f'Let op: {datum} is geen maandag.')
    k = kalender(maandag)

    slides = deck(inv, k)
    os.makedirs(os.path.join(map_, 'preview'), exist_ok=True)
    for i, (els, dark) in enumerate(slides, 1):
        if keuze and i not in keuze:
            continue
        # De HTML is alleen tussenstap voor de screenshot, die hoort niet in de repo.
        html = render_html(els, os.path.join(tempfile.gettempdir(), f'sales-update-slide-{i}.html'), dark=dark)
        shoot(html, os.path.join(map_, 'preview', f'slide-{i}.png'))
        print('preview', i)

    if not keuze:
        pptx = os.path.join(map_, f'sales-update-{datum}-bewerkbaar.pptx')
        render_pptx([els for els, _ in slides], pptx)
        controle(pptx)
        print('pptx', pptx)


def controle(pad):
    """Niets mag buiten het canvas vallen (pitch-deck SKILL.md)."""
    from pptx import Presentation
    p = Presentation(pad)
    for i, sl in enumerate(p.slides, 1):
        buiten = [s.shape_id for s in sl.shapes
                  if s.left < 0 or s.top < 0
                  or s.left + s.width > p.slide_width or s.top + s.height > p.slide_height]
        if buiten:
            print(f'slide {i}: buiten canvas {buiten}')


if __name__ == '__main__':
    main()
